"""Quickstart scenarios 1-6 validation script."""
import asyncio, json, sys, tempfile
from docx import Document
from pathlib import Path

sys.path.insert(0, "src")
sys.path.insert(0, "tests")
sys.path.insert(0, ".")

from agentic_rag_mcp.config import Config
from agentic_rag_mcp.server import build_server

def call(server, name, args):
    raw = asyncio.run(server.call_tool(name, args))
    structured = [b["result"] for b in raw if isinstance(b, dict) and "result" in b]
    if structured:
        return structured[0]
    texts = [getattr(b, "text", "") for b in raw if getattr(b, "text", "")]
    joined = "".join(texts).strip()
    return json.loads(joined) if joined else None

# build fixture corpus
import subprocess
subprocess.run([".venv/bin/python", "-c", """
from docx import Document; from pptx import Presentation; import openpyxl
from pathlib import Path
d = Path('/tmp/qs-validate/docs'); d.mkdir(parents=True, exist_ok=True)
doc = Document(); doc.add_paragraph('The Eiffel Tower is 330 metres tall, located in Paris.')
doc.save(d/'paris.docx')
wb = openpyxl.Workbook(); ws = wb.active
ws.append(['metric','value']); ws.append(['uptime', 99.9])
wb.save(d/'metrics.xlsx')
prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Roadmap'; s.placeholders[1].text = 'Q3 goal: ship the MCP server.'
prs.save(d/'roadmap.pptx')
(d/'broken.pdf').write_bytes(b'not a real pdf')
import sqlite3
c = sqlite3.connect('/tmp/qs-validate/data.sqlite3')
c.execute('CREATE TABLE facts (id INTEGER PRIMARY KEY, body TEXT)')
c.execute("INSERT INTO facts (body) VALUES ('The Pacific is the largest ocean.')")
c.commit(); c.close()
print('corpus ready')
"""], check=True)

tmp = Path(tempfile.mkdtemp())
cfg = Config(state_db=tmp/"state.sqlite3", vector_dir=tmp/"vector",
             laya_url="http://127.0.0.1:1", max_iterations=2)
srv = build_server(cfg)

results = {}

# Scenario 1: local folder -> fast answer
out = call(srv, "sources_add", {"type": "local_folder",
                                "config": {"path": "/tmp/qs-validate/docs"}})
s1_ok_corpus = out["sync"]["added"] == 3 and len(out["sync"]["skipped"]) == 1
ans = call(srv, "ask", {"question": "How tall is the Eiffel Tower?", "mode": "fast"})
s1 = s1_ok_corpus and "330" in ans["answer"] and ans["citations"]
results["S1 fast path + citations"] = bool(s1)

# Scenario 2: incremental update + delete
d2 = Document(); d2.add_paragraph("The Amazon is the largest river by discharge.")
d2.save("/tmp/qs-validate/docs/update.docx")
st = call(srv, "sources_sync", {"source_id": out["source_id"]})
s2a = st[out["source_id"]]["added"] >= 1
Path("/tmp/qs-validate/docs/roadmap.pptx").unlink()
st = call(srv, "sources_sync", {"source_id": out["source_id"]})
s2b = st[out["source_id"]]["deleted"] >= 1
res = call(srv, "search", {"query": "Q3 goal ship MCP server"})
s2c = all("roadmap.pptx" != h["title"] for h in res)
results["S2 incremental add/delete"] = bool(s2a and s2b and s2c)

# Scenario 3: multi-source deep research
sid2 = call(srv, "sources_add", {"type": "sqlite",
    "config": {"db_path": "/tmp/qs-validate/data.sqlite3", "table": "facts"}})["source_id"]
ans3 = call(srv, "ask", {"question":
    "Compare the Eiffel Tower height and the largest ocean in the facts table",
    "mode": "deep"})
s3 = ans3["classification"] == "deliberative" and len(ans3["trace"]) >= 1
results["S3 deep research trace"] = bool(s3)

# Scenario 4: web search fallback (no provider configured -> degraded reported, local answer still returned)
ans4 = call(srv, "ask", {"question": "latest news about python release today", "mode": "deep"})
s4 = isinstance(ans4.get("answer"), str) and ans4.get("citations", None) is not None
results["S4 web fallback graceful"] = bool(s4)

# Scenario 5: error handling
s5a = out["sync"]["skipped"] and "broken.pdf" in out["sync"]["skipped"][0]["locator"]
fresh = build_server(Config(state_db=Path(tempfile.mkdtemp())/"s.sqlite3",
                            vector_dir=Path(tempfile.mkdtemp())/"v",
                            laya_url="http://127.0.0.1:1", max_iterations=1))
ans5 = call(fresh, "ask", {"question": "What is invisible to all sources xyz 123?",
                           "mode": "deep", "max_iterations": 1})
s5b = ans5["confidence"] == "low" and ans5["sufficiency"] == "exhausted"
st5 = call(fresh, "status", {})
s5c = st5["decision_layer"] == "fallback"
results["S5 error paths (skip/low-confidence/laya-fallback)"] = bool(s5a and s5b and s5c)

# Scenario 6: tool surface check
tools = asyncio.run(srv.list_tools())
names = {t.name for t in tools}
need = {"ask", "search", "sources_add", "sources_list", "sources_remove", "sources_sync",
        "graph_query", "status"}
s6 = need.issubset(names)
missing = need - names
results[f"S6 MCP tool surface (missing: {sorted(missing) if missing else 'none'})"] = bool(s6)

print()
allok = True
for k, v in results.items():
    print(f"  {'PASS' if v else 'FAIL'}: {k}")
    allok &= v
print()
print("ALL PASS" if allok else "SOME FAILED")
sys.exit(0 if allok else 1)
