"""Document and audio extraction.

Corrupt/encrypted/no-speech inputs raise ParseSkipped(reason) — the ingestion
pipeline records status=failed|skipped and continues with the next item.
"""

from __future__ import annotations

from pathlib import Path

from ..errors import IngestionError


class ParseSkipped(IngestionError):
    pass


_EXTS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".pptx": "pptx"}
_AUDIO_EXTS = {".mp3", ".wav", ".m4a"}


def media_type_for(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in _EXTS:
        return _EXTS[ext]
    if ext in _AUDIO_EXTS:
        return "audio"
    raise ParseSkipped(f"unsupported file type: {ext or '(none)'}")


def extract_text(path: Path) -> str:
    mt = media_type_for(path)
    try:
        if mt == "pdf":
            return _extract_pdf(path)
        if mt == "docx":
            return _extract_docx(path)
        if mt == "xlsx":
            return _extract_xlsx(path)
        if mt == "pptx":
            return _extract_pptx(path)
    except IngestionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseSkipped(f"{path.name}: {exc.__class__.__name__}: {exc}") from exc
    return transcribe(path)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ParseSkipped(f"corrupt or unreadable pdf: {path.name}") from exc
    if reader.is_encrypted:
        raise ParseSkipped(f"encrypted pdf: {path.name}")
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if not text:
        raise ParseSkipped(f"no extractable text: {path.name}")
    return text


def _extract_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text).strip()
    if not text:
        raise ParseSkipped(f"no extractable text: {path.name}")
    return text


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                lines.append(" | ".join(vals))
    text = "\n".join(lines).strip()
    if not text:
        raise ParseSkipped(f"no extractable text: {path.name}")
    return text


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text for p in shape.text_frame.paragraphs if p.text)
    text = "\n".join(lines).strip()
    if not text:
        raise ParseSkipped(f"no extractable text: {path.name}")
    return text


def transcribe(path: Path, model_size: str = "base") -> str:
    """Local faster-whisper transcription (optional [audio] extra)."""
    try:
        from faster_whisper import WhisperModel
    except ImportError as exc:
        raise ParseSkipped("audio support not installed (pip install .[audio])") from exc
    model = WhisperModel(model_size, device="cpu", compute_type="int8")
    try:
        segments, _info = model.transcribe(str(path))
    except TypeError as exc:
        # faster-whisper + new PyAV: metadata_errors kwarg removed upstream
        raise ParseSkipped(f"audio decode unsupported: {path.name} ({exc})") from exc
    except Exception as exc:
        raise ParseSkipped(f"audio unreadable: {path.name} ({exc.__class__.__name__})") from exc
    text = " ".join(seg.text for seg in segments).strip()
    if not text:
        raise ParseSkipped(f"no speech detected: {path.name}")
    return text
